#!/usr/bin/env python3
r"""Import a day's sales sheet into Adminator (receipts + line items).

Reads ``pastry_shop_sales_*.xlsx`` ("Sales Lines" sheet) and records one sale
per order via the REST API (``POST /sales/sales/record/``). Lines are grouped by
**Order Id** (and split by payment method if an order ever mixes them). For each
line the product comes from **Product Name** and the count from **Qty**. Lines
are priced at the product's **catalog selling price** by default (the original
price from the app); pass ``--use-sheet-prices`` to price them at **Total Price ÷
Qty** from the sheet instead. The payment method picks the wallet the takings
land in:

    MoMo  → "Mobile money balance"      Cash → "Petite caisse"      Card → "Bank"

Edge case handled
-----------------
*If a product is out of stock* when the sale is recorded (the API refuses to let
stock go negative), the script bakes what's missing — it runs production for the
shortfall and retries the sale. Those production runs go through the same path as
``import_production.py``: *if production itself is short a raw material*, the
material is topped up and the run retried. So an out-of-stock sale cascades into
"produce the product (buying any raw materials it needs), then sell it".

Voided lines are skipped. Product names are matched to the catalog case/accent-
insensitively with a small alias table; lines whose product can't be matched
confidently are reported and skipped (the rest of the order still goes in).

Safe by default: **dry-run** (parse + match + plan, no writes). Pass ``--commit``
to record the sales.

Usage
-----
    python scripts/import_sales.py                 # preview only (catalog prices)
    python scripts/import_sales.py --commit        # actually record
    python scripts/import_sales.py --commit --use-sheet-prices   # price from the sheet
    python scripts/import_sales.py --commit --limit 1   # just the first order

Requirements (host): ``pip install openpyxl`` (HTTP uses the standard library).
"""
from __future__ import annotations

import argparse
import datetime
import json
import math
import os
import re
import sys
import unicodedata
import urllib.error
import urllib.request
from collections import OrderedDict
from decimal import ROUND_CEILING, Decimal
from difflib import SequenceMatcher

try:
    import openpyxl
except ModuleNotFoundError:
    sys.exit('openpyxl is not installed.\n  Fix: python -m pip install openpyxl')

# ── defaults (override on the CLI or via env) ────────────────────────────────
DEFAULT_BASE_URL = os.environ.get("ADMINATOR_API", "http://localhost:8001/api/v1")
DEFAULT_EMAIL = os.environ.get("ADMINATOR_EMAIL", "admin@adminator.local")
DEFAULT_PASSWORD = os.environ.get("ADMINATOR_PASSWORD", "admin12345")
DEFAULT_FILE = os.environ.get(
    "SALES_XLSX",
    r"D:\PC DISAINE\toky\Perso-D\red\sales\2026_06_16\pastry_shop_sales_2026-06-16.xlsx",
)
DEFAULT_DATE = "2026-06-16"        # fallback if a row has no Date
DEFAULT_OFFSET = "+02:00"          # Africa/Kigali (no DST)
SHEET_NAME = "Sales Lines"

# Payment method (sheet) -> API value, and the wallet each one settles into.
PAYMENT_METHOD = {"momo": "mobile_money", "cash": "cash", "card": "card"}
WALLET_FOR_METHOD = {
    "mobile_money": "Mobile money balance",
    "cash": "Petite caisse",
    "card": "Bank",
}

# Curated aliases: sales-sheet name (normalized) -> exact catalog name.
# High-confidence only: word order, typos, and sale annotations like "(Promo)".
PRODUCT_ALIASES = {
    "baguette brown": "Brown baguette",
    "coffee latte": "Latte",
    "panini tuna": "Tuna panini",
    "salad chicken": "Chicken salad",
    "sourdhough bread promo": "Sourdough bread",
}

# Judgment-call mappings — DISABLED by default. These sheet names have no exact
# catalog product, so the script reports them as "unmatched" and skips them
# (the rest of the order still imports). Review each, and move any you accept up
# into PRODUCT_ALIASES before committing. Left here so the decision is explicit.
#   "coca":              "Coca cola",   # plain "Coca": Coca cola? Coca zero?
#   "pizza":             "Pizza roll",  # only pizza-like product (sheet 13000 vs catalog 6000)
#   "croissant fromage": "Ham & cheese croissant",  # cheese croissant; catalog only has ham & cheese
#   "chips":             <none>         # no catalog product — would need creating
REVIEW_ALIASES: dict[str, str] = {}

MATCH_THRESHOLD = 0.82
MAX_TOPUP_ROUNDS = 60       # raw-material top-up retries within one production run
MAX_PRODUCE_ROUNDS = 40     # produce-then-retry passes for one out-of-stock sale

INSUFFICIENT_RE = re.compile(
    r"Not enough stock of (.+?): have ([\d.]+), requested ([\d.]+)\."
)


# ── small helpers ────────────────────────────────────────────────────────────
def normalize(name: str) -> str:
    """Lowercase, strip accents, reduce to space-separated alphanumeric tokens.

    Parenthetical content is kept so "Cake" and "Cake (6 persons)" stay distinct.
    """
    if name is None:
        return ""
    s = unicodedata.normalize("NFKD", str(name))
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.lower()
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def parse_insufficient(detail: str | None):
    if not detail:
        return None
    m = INSUFFICIENT_RE.search(detail)
    if not m:
        return None
    return m.group(1), Decimal(m.group(2)), Decimal(m.group(3))


def topup_amount(have: Decimal, requested: Decimal) -> Decimal:
    deficit = Decimal(requested) - Decimal(have)
    if deficit <= 0:
        deficit = Decimal("0.0001")
    return deficit.quantize(Decimal("0.0001"), rounding=ROUND_CEILING)


def to_iso(date_val, time_val, offset: str, fallback_date: str) -> str:
    """Build an ISO8601 timestamp with explicit offset from sheet Date + Time."""
    if isinstance(date_val, (datetime.date, datetime.datetime)):
        d = date_val.strftime("%Y-%m-%d")
    elif date_val:
        d = str(date_val).strip()[:10]
    else:
        d = fallback_date
    if isinstance(time_val, (datetime.time, datetime.datetime)):
        t = time_val.strftime("%H:%M:%S")
    else:
        t = (str(time_val).strip() if time_val else "12:00")
        parts = t.split(":")
        if len(parts) == 2:
            t = f"{int(parts[0]):02d}:{int(parts[1]):02d}:00"
    return f"{d}T{t}{offset}"


# ── API client ───────────────────────────────────────────────────────────────
class ApiError(Exception):
    def __init__(self, status: int, body):
        self.status = status
        self.body = body
        super().__init__(f"HTTP {status}: {body}")

    @property
    def detail(self) -> str:
        if isinstance(self.body, dict):
            d = self.body.get("detail")
            return str(d) if d else json.dumps(self.body)
        return str(self.body)


class Api:
    def __init__(self, base_url: str, email: str, password: str):
        self.base = base_url.rstrip("/")
        self.token: str | None = None
        _, data = self._request(
            "POST", "/auth/login/", {"email": email, "password": password}, auth=False
        )
        self.token = data["access"]

    def _request(self, method, path, payload=None, auth=True):
        url = self.base + path
        body = json.dumps(payload).encode() if payload is not None else None
        req = urllib.request.Request(url, data=body, method=method)
        req.add_header("Content-Type", "application/json")
        if auth and self.token:
            req.add_header("Authorization", f"Bearer {self.token}")
        try:
            with urllib.request.urlopen(req, timeout=60) as resp:
                raw = resp.read().decode()
                return resp.status, (json.loads(raw) if raw else None)
        except urllib.error.HTTPError as exc:
            raw = exc.read().decode()
            try:
                parsed = json.loads(raw)
            except ValueError:
                parsed = {"detail": raw[:300]}
            raise ApiError(exc.code, parsed)
        except urllib.error.URLError as exc:
            sys.exit(f"Could not reach {url}: {exc}. Is the scratch stack up?")

    def get_all(self, path: str) -> list:
        sep = "&" if "?" in path else "?"
        _, data = self._request("GET", f"{path}{sep}page_size=1000")
        if isinstance(data, dict) and "results" in data:
            rows = list(data["results"])
            while data.get("next"):
                nxt = data["next"].split("/api/v1", 1)[-1]
                _, data = self._request("GET", nxt)
                rows.extend(data["results"])
            return rows
        return data or []

    def record_sale(self, payload: dict) -> dict:
        _, data = self._request("POST", "/sales/sales/record/", payload)
        return data

    def execute_production(self, product_id, quantity, scheduled_for, notes=""):
        _, data = self._request(
            "POST", "/production/runs/execute/",
            {"product": product_id, "quantity": str(quantity),
             "scheduled_for": scheduled_for, "status": "completed", "notes": notes},
        )
        return data

    def receive_material(self, raw_material_id, quantity, note="", reference=""):
        _, data = self._request(
            "POST", "/inventory/stock/receive/",
            {"raw_material": raw_material_id, "quantity": str(quantity),
             "note": note, "reference": reference},
        )
        return data


# ── name matching ────────────────────────────────────────────────────────────
class Matcher:
    def __init__(self, products: list[dict], aliases: dict[str, str]):
        self.products = products
        self.by_norm: dict[str, dict] = {}
        for p in products:
            self.by_norm.setdefault(normalize(p["name"]), p)
        self.alias_norm: dict[str, dict] = {}
        for sheet_name, catalog_name in aliases.items():
            target = self.by_norm.get(normalize(catalog_name))
            if target is None:
                print(f"  ! alias target not in catalog: {catalog_name!r}", file=sys.stderr)
            else:
                self.alias_norm[normalize(sheet_name)] = target

    def match(self, name: str):
        n = normalize(name)
        if not n:
            return None, "empty", 0.0
        if n in self.alias_norm:
            return self.alias_norm[n], "alias", 1.0
        if n in self.by_norm:
            return self.by_norm[n], "exact", 1.0
        best, best_score = None, 0.0
        n_sorted = " ".join(sorted(n.split()))
        for p in self.products:
            pn = normalize(p["name"])
            score = max(
                SequenceMatcher(None, n, pn).ratio(),
                SequenceMatcher(None, n_sorted, " ".join(sorted(pn.split()))).ratio(),
            )
            if score > best_score:
                best, best_score = p, score
        if best_score >= MATCH_THRESHOLD:
            return best, "fuzzy", best_score
        return None, "unmatched", best_score


def index_materials(stock_rows: list[dict]) -> dict[str, list[dict]]:
    idx: dict[str, list[dict]] = {}
    for row in stock_rows:
        if row.get("kind") != "raw_material" or not row.get("raw_material"):
            continue
        entry = {"id": row["raw_material"], "name": row["item_name"],
                 "stock": Decimal(str(row["quantity"]))}
        idx.setdefault(normalize(row["item_name"]), []).append(entry)
    return idx


def resolve_material(materials: dict, name: str, have: Decimal):
    cands = materials.get(normalize(name), [])
    if not cands:
        return None
    if len(cands) == 1:
        return cands[0]
    for c in cands:
        if c["stock"] == have:
            return c
    return cands[0]


# ── production helpers (shared shape with import_production.py) ───────────────
def produce_with_topup(api: Api, product: dict, quantity: Decimal, date: str,
                       materials: dict, log) -> dict:
    """Record a production run, topping up any short raw material and retrying."""
    notes = f"Auto-produced for sales import ({date})"
    last_short = None
    for _ in range(MAX_TOPUP_ROUNDS):
        try:
            return api.execute_production(product["id"], quantity, date, notes)
        except ApiError as exc:
            if exc.status != 400:
                raise
            info = parse_insufficient(exc.detail)
            if info is None:
                raise
            mat_name, have, requested = info
            mat = resolve_material(materials, mat_name, have)
            if mat is None:
                raise RuntimeError(f"raw material {mat_name!r} (from error) not in catalog")
            amount = topup_amount(have, requested)
            api.receive_material(
                mat["id"], amount,
                note=f"Auto top-up so '{product['name']}' can be produced (sales import)",
                reference="IMPORT-SALE",
            )
            mat["stock"] = have + amount
            log(f"        ↑ topped up {mat_name}: +{amount}")
            last_short = mat_name
    raise RuntimeError(f"gave up topping up materials (stuck on {last_short})")


def record_sale_with_production(api: Api, payload: dict, matcher: Matcher,
                                materials: dict, prod_date: str, log) -> dict:
    """Record a sale; bake any out-of-stock product on demand, then retry."""
    for _ in range(MAX_PRODUCE_ROUNDS):
        try:
            return api.record_sale(payload)
        except ApiError as exc:
            if exc.status != 400:
                raise
            info = parse_insufficient(exc.detail)
            if info is None:
                raise                      # genuine validation error — surface it
            pname, have, requested = info
            product, _, _ = matcher.match(pname)
            if product is None:
                raise RuntimeError(f"out-of-stock product {pname!r} not in catalog")
            qty = Decimal(max(1, math.ceil(Decimal(requested) - Decimal(have))))
            log(f"      ⚙ out of stock: {pname} (have {have}, need {requested}) → producing {qty}")
            produce_with_topup(api, product, qty, prod_date, materials, log)
    raise RuntimeError("too many produce/retry passes for one sale")


# ── spreadsheet → grouped orders ─────────────────────────────────────────────
def read_orders(path: str, fallback_date: str, offset: str):
    """Return OrderedDict keyed by (order_id, method) -> {meta, lines:[...]}."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    col = {name: i for i, name in enumerate(header)}

    def cell(r, name):
        i = col.get(name)
        return r[i] if i is not None and i < len(r) else None

    orders: "OrderedDict[tuple, dict]" = OrderedDict()
    voided = 0
    for r in rows[1:]:
        name = cell(r, "Product Name")
        order_id = cell(r, "Order Id")
        if name is None or order_id is None:
            continue
        if str(cell(r, "Status")).strip().lower() == "voided":
            voided += 1
            continue
        method_raw = str(cell(r, "Payment Method") or "").strip().lower()
        method = PAYMENT_METHOD.get(method_raw)
        key = (order_id, method or method_raw)
        if key not in orders:
            orders[key] = {
                "order_id": order_id,
                "method_raw": method_raw,
                "method": method,
                "occurred_at": to_iso(cell(r, "Date"), cell(r, "Time"), offset, fallback_date),
                "lines": [],
            }
        orders[key]["lines"].append({
            "product_name": str(name).strip(),
            "qty": cell(r, "Qty"),
            "total": cell(r, "Total Price"),
        })
    return orders, voided


# ── main ─────────────────────────────────────────────────────────────────────
def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass

    p = argparse.ArgumentParser(description="Import daily sales from the pastry sales xlsx.")
    p.add_argument("--file", default=DEFAULT_FILE)
    p.add_argument("--base-url", default=DEFAULT_BASE_URL)
    p.add_argument("--email", default=DEFAULT_EMAIL)
    p.add_argument("--password", default=DEFAULT_PASSWORD)
    p.add_argument("--date", default=DEFAULT_DATE, help="Fallback date + production scheduled_for date.")
    p.add_argument("--utc-offset", default=DEFAULT_OFFSET, help="Timezone offset for occurred_at (default Africa/Kigali +02:00).")
    p.add_argument("--use-sheet-prices", action="store_true",
                   help="Price each line at Total Price / Qty from the sheet. "
                        "Default: use the product's catalog selling price from the app.")
    p.add_argument("--commit", action="store_true", help="Actually write (default: dry-run preview).")
    p.add_argument("--limit", type=int, default=0, help="Only process the first N orders (testing).")
    args = p.parse_args()

    mode = "COMMIT" if args.commit else "DRY-RUN (no writes)"
    pricing = "sheet (Total / Qty)" if args.use_sheet_prices else "app catalog selling price"
    print(f"Sales import · {mode}")
    print(f"File   : {args.file}")
    print(f"Pricing: {pricing}")
    print(f"Target : {args.base_url}\n")

    orders, voided = read_orders(args.file, args.date, args.utc_offset)

    api = Api(args.base_url, args.email, args.password)
    products = api.get_all("/catalog/products/")
    materials = index_materials(api.get_all("/inventory/stock/?kind=raw_material"))
    matcher = Matcher(products, PRODUCT_ALIASES)
    wallets = {w["name"]: w["id"] for w in api.get_all("/finance/wallets/")}
    for needed in WALLET_FOR_METHOD.values():
        if needed not in wallets:
            sys.exit(f"Required wallet {needed!r} not found. Have: {list(wallets)}")
    print(f"Loaded {len(products)} products, {len(wallets)} wallets. "
          f"{len(orders)} order/method groups ({voided} voided line(s) skipped).\n")

    # Build sale payloads, resolving products + wallets.
    sales, unmatched_lines, bad_orders = [], [], []
    by_method_total = {"mobile_money": Decimal("0"), "cash": Decimal("0"), "card": Decimal("0")}

    for (order_id, _), o in orders.items():
        if o["method"] is None:
            bad_orders.append((order_id, f"unknown payment method {o['method_raw']!r}"))
            continue
        items, line_descr = [], []
        order_total = Decimal("0")
        for ln in o["lines"]:
            product, method, score = matcher.match(ln["product_name"])
            if product is None:
                unmatched_lines.append((order_id, ln["product_name"], score))
                continue
            qty = Decimal(str(ln["qty"]))
            item = {"product": product["id"], "quantity": str(qty)}
            if args.use_sheet_prices:
                sheet_total = Decimal(str(ln["total"]))
                unit_price = (sheet_total / qty).quantize(Decimal("0.01")) if qty else Decimal("0")
                item["unit_price"] = str(unit_price)
            else:
                # Omit unit_price → the API falls back to product.selling_price,
                # i.e. the original catalog price from the app.
                unit_price = Decimal(str(product["selling_price"]))
            items.append(item)
            order_total += (unit_price * qty).quantize(Decimal("0.01"))
            tag = method if method != "fuzzy" else f"~{score:.2f}"
            line_descr.append(f"{product['name']}×{qty}@{unit_price}[{tag}]")
        if not items:
            bad_orders.append((order_id, "no matchable items"))
            continue
        payload = {
            "items": items,
            "payment_method": o["method"],
            "channel": "counter",
            "occurred_at": o["occurred_at"],
            "wallet": wallets[WALLET_FOR_METHOD[o["method"]]],
            "notes": f"Imported from sales sheet — order {order_id}",
        }
        sales.append({"order_id": order_id, "method": o["method"], "total": order_total,
                      "payload": payload, "descr": line_descr})
        by_method_total[o["method"]] += order_total
        print(f"  order {str(order_id):>3} [{o['method']:<12}] {WALLET_FOR_METHOD[o['method']]:<20} "
              f"{order_total:>8} : " + ", ".join(line_descr))

    if unmatched_lines:
        print("\n  Unmatched lines (skipped):")
        for oid, nm, sc in unmatched_lines:
            print(f"    order {oid}: {nm!r}  (best score {sc:.2f})")

    if args.limit:
        sales = sales[: args.limit]

    done, failed = [], []
    if args.commit and sales:
        print("\nRecording sales…")
        for s in sales:
            try:
                sale = record_sale_with_production(
                    api, s["payload"], matcher, materials, args.date, log=print)
                done.append(sale)
                print(f"  ✓ order {s['order_id']} [{s['method']}] {s['total']} "
                      f"→ {sale['receipt_number']} (total {sale['total']})")
            except (ApiError, RuntimeError) as exc:
                msg = exc.detail if isinstance(exc, ApiError) else str(exc)
                failed.append((s["order_id"], msg))
                print(f"  ✗ order {s['order_id']} [{s['method']}] — {msg}")

    # ── summary ──
    print("\n" + "=" * 64)
    print("Summary")
    print(f"  sales planned        : {len(sales)}")
    print(f"  voided lines skipped : {voided}")
    print(f"  unmatched lines      : {len(unmatched_lines)}")
    print(f"  orders skipped       : {len(bad_orders)}"
          + (f"  → {bad_orders}" if bad_orders else ""))
    print("  planned totals by method (matched lines only):")
    for m, label in (("mobile_money", "MoMo"), ("cash", "Cash"), ("card", "Card")):
        print(f"      {label:<5} {by_method_total[m]:>10}")
    print(f"      {'TOTAL':<5} {sum(by_method_total.values()):>10}")
    if args.commit:
        print(f"  sales recorded       : {len(done)}")
        print(f"  sales failed         : {len(failed)}")
    else:
        print("\n  DRY-RUN — nothing written. Re-run with --commit to record these sales.")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
