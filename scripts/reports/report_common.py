#!/usr/bin/env python3
r"""Shared plumbing for the Adminator reporting scripts.

The three report scripts — ``daily_report.py``, ``weekly_report.py`` and
``monthly_report.py`` — all read from the same REST API, bucket the same rows
by date, and format money / percentages the same way. That common machinery
lives here so each report file can stay focused on *what* it measures rather
than *how* it talks to the backend.

What this module provides
-------------------------
* :class:`Api` — a tiny stdlib-only (``urllib``) JWT client, the same shape the
  import scripts use: ``POST /auth/login/`` for a bearer token, then
  ``get_all()`` which follows DRF pagination.
* :class:`DataStore` — fetches each collection (sales, production runs,
  products, expenses, transactions, wallets) at most once and caches it, so a
  report can slice the same data several ways for free.
* Date helpers (:func:`sale_date`, :func:`in_range`, :func:`month_bounds`,
  :func:`week_bounds`) that bucket rows by the shop's *local* calendar day. The
  API returns timestamps with an explicit ``+02:00`` offset (Africa/Kigali), so
  the date portion of the string already is the local day — no conversion
  needed.
* Aggregation helpers shared by more than one report
  (:func:`product_sales`, :func:`payment_breakdown`, :func:`product_production`,
  …) that take an already-filtered list of rows and return plain dicts.
* Formatting + a small table/section renderer for consistent, readable output,
  plus ``--json`` support via :func:`emit`.

Config comes from the same environment variables the import scripts honour
(``ADMINATOR_API`` / ``ADMINATOR_EMAIL`` / ``ADMINATOR_PASSWORD``), plus a few
reporting-only ones documented in :data:`ENV_HELP`. Nothing here writes — the
reports are strictly read-only.
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import urllib.error
import urllib.request
from calendar import monthrange
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

# ── config (override on the CLI or via env) ──────────────────────────────────
DEFAULT_BASE_URL = os.environ.get("ADMINATOR_API", "http://localhost:8001/api/v1")
DEFAULT_EMAIL = os.environ.get("ADMINATOR_EMAIL", "admin@adminator.local")
DEFAULT_PASSWORD = os.environ.get("ADMINATOR_PASSWORD", "admin12345")
CURRENCY = os.environ.get("ADMINATOR_CURRENCY", "RWF")

ENV_HELP = """\
Environment variables honoured by the report scripts:
  ADMINATOR_API                 Base API URL (default http://localhost:8001/api/v1)
  ADMINATOR_EMAIL               Login email (default admin@adminator.local)
  ADMINATOR_PASSWORD            Login password
  ADMINATOR_CURRENCY            Currency code shown next to amounts (default RWF)
  ADMINATOR_LABOR_COST_MONTHLY  Monthly labour cost — labour is not tracked in
                                the app, so the monthly P&L reads it from here.
  ADMINATOR_LABOR_HOURS_WEEKLY  Staff hours worked in a week — feeds the weekly
                                "employee hours per 1,000 revenue" KPI.
  ADMINATOR_COMPLAINTS_WEEKLY   Customer complaints in a week (not tracked in app).
  ADMINATOR_DELIVERY_DELAYS_WEEKLY  Late deliveries in a week (not tracked in app).
"""

ZERO = Decimal("0")
CENTS = Decimal("0.01")

#: Name of the wallet that backs the physical till. Cash-drawer reconciliation
#: only cares about expenses paid out of *this* wallet, not every "cash"
#: payment_method row (an expense can be tagged cash but paid from elsewhere,
#: or vice versa) — see import_sales.py's WALLET_FOR_METHOD mapping.
CASH_DRAWER_WALLET = "Petite caisse"


# ── numeric helpers ──────────────────────────────────────────────────────────
def D(value) -> Decimal:
    """Coerce anything money-ish (str/int/float/Decimal/None) to a Decimal."""
    if value is None or value == "":
        return ZERO
    if isinstance(value, Decimal):
        return value
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return ZERO


def q2(value: Decimal) -> Decimal:
    """Round to 2 dp (half-up), the precision the API stores money at."""
    return D(value).quantize(CENTS, rounding=ROUND_HALF_UP)


def safe_div(numer, denom):
    """``numer / denom`` as a Decimal, or ``None`` when ``denom`` is zero."""
    numer, denom = D(numer), D(denom)
    if denom == 0:
        return None
    return numer / denom


def pct(part, whole):
    """``part / whole`` expressed as a percentage (a number), or ``None``."""
    ratio = safe_div(part, whole)
    return None if ratio is None else ratio * 100


# ── formatting ───────────────────────────────────────────────────────────────
def money(value, currency: str = CURRENCY) -> str:
    """Format an amount with thousands separators and the currency code."""
    return f"{q2(value):,.2f} {currency}"


def signed_money(value, currency: str = CURRENCY) -> str:
    """Like :func:`money` but always shows an explicit +/- sign."""
    v = q2(value)
    return f"{'+' if v >= 0 else '-'}{abs(v):,.2f} {currency}"


def percent(value, decimals: int = 1) -> str:
    """Format a percentage number, or ``n/a`` when ``value`` is ``None``."""
    if value is None:
        return "n/a"
    return f"{D(value):,.{decimals}f}%"


def num(value, decimals: int = 0) -> str:
    """Format a plain count (receipts, runs, complaints) with thousands separators."""
    if value is None:
        return "n/a"
    return f"{D(value):,.{decimals}f}"


def qty(value, max_decimals: int = 2) -> str:
    """Format a unit quantity: whole numbers show no decimals, fractional ones
    keep up to ``max_decimals`` (trailing zeros trimmed). Production and sales
    quantities can be fractional, so plain ``num`` would hide small runs."""
    if value is None:
        return "n/a"
    d = D(value)
    if d == d.to_integral_value():
        return f"{d:,.0f}"
    return f"{d:,.{max_decimals}f}".rstrip("0").rstrip(".")


# ── API client (stdlib urllib; JWT bearer) ───────────────────────────────────
class ApiError(Exception):
    def __init__(self, status: int, body):
        self.status = status
        self.body = body
        super().__init__(f"HTTP {status}: {body}")


class Api:
    """Minimal read client: log in for a token, then GET (with pagination)."""

    def __init__(self, base_url: str, email: str, password: str):
        self.base = base_url.rstrip("/")
        self.token: str | None = None
        _, data = self._request(
            "POST", "/auth/login/", {"email": email, "password": password}, auth=False
        )
        self.token = data["access"]

    def _request(self, method, path, payload=None, auth=True):
        url = self.base + path
        body = json.dumps(payload).encode() if payload is not None else None
        req = urllib.request.Request(url, data=body, method=method)
        req.add_header("Content-Type", "application/json")
        if auth and self.token:
            req.add_header("Authorization", f"Bearer {self.token}")
        try:
            with urllib.request.urlopen(req, timeout=60) as resp:
                raw = resp.read().decode()
                return resp.status, (json.loads(raw) if raw else None)
        except urllib.error.HTTPError as exc:
            raw = exc.read().decode()
            try:
                parsed = json.loads(raw)
            except ValueError:
                parsed = {"detail": raw[:300]}
            raise ApiError(exc.code, parsed)
        except urllib.error.URLError as exc:
            sys.exit(f"Could not reach {url}: {exc}. Is the stack up?")

    def get_all(self, path: str) -> list:
        """Follow DRF pagination and return every result row."""
        sep = "&" if "?" in path else "?"
        _, data = self._request("GET", f"{path}{sep}page_size=1000")
        if isinstance(data, dict) and "results" in data:
            rows = list(data["results"])
            while data.get("next"):
                nxt = data["next"].split("/api/v1", 1)[-1]
                _, data = self._request("GET", nxt)
                rows.extend(data["results"])
            return rows
        return data or []


class DataStore:
    """Fetch-once, slice-many cache over the read endpoints a report needs."""

    def __init__(self, api: Api):
        self.api = api
        self._cache: dict[str, list] = {}

    def _get(self, key: str, path: str) -> list:
        if key not in self._cache:
            self._cache[key] = self.api.get_all(path)
        return self._cache[key]

    def sales(self) -> list:
        return self._get("sales", "/sales/sales/")

    def runs(self) -> list:
        return self._get("runs", "/production/runs/")

    def products(self) -> list:
        return self._get("products", "/catalog/products/")

    def expenses(self) -> list:
        return self._get("expenses", "/finance/expenses/")

    def transactions(self) -> list:
        return self._get("transactions", "/finance/transactions/")

    def wallets(self) -> list:
        return self._get("wallets", "/finance/wallets/")


# ── date helpers ─────────────────────────────────────────────────────────────
def parse_date(value: str) -> date:
    return date.fromisoformat(value)


def sale_date(sale: dict) -> date:
    """Local calendar day of a sale (the ``occurred_at`` offset is shop-local)."""
    return date.fromisoformat(sale["occurred_at"][:10])


def in_range(d: date, start: date, end: date) -> bool:
    """Inclusive ``start <= d <= end``."""
    return start <= d <= end


def month_bounds(year: int, month: int) -> tuple[date, date]:
    last = monthrange(year, month)[1]
    return date(year, month, 1), date(year, month, last)


def parse_month(value: str) -> tuple[date, date]:
    """``"2026-06"`` (or ``"2026-06-15"``) -> (first_day, last_day) of that month."""
    parts = value.split("-")
    year, month = int(parts[0]), int(parts[1])
    return month_bounds(year, month)


def week_bounds(ending: date, length: int = 7) -> tuple[date, date]:
    """The ``length``-day window ending on (and including) ``ending``."""
    return ending - timedelta(days=length - 1), ending


# ── filtered selectors ───────────────────────────────────────────────────────
def sales_between(sales: list, start: date, end: date) -> list:
    return [s for s in sales if in_range(sale_date(s), start, end)]


def completed_runs_between(runs: list, start: date, end: date) -> list:
    out = []
    for r in runs:
        if r.get("status") != "completed":
            continue
        d = date.fromisoformat(r["scheduled_for"])
        if in_range(d, start, end):
            out.append(r)
    return out


def expenses_between(expenses: list, start: date, end: date) -> list:
    return [
        e for e in expenses
        if in_range(date.fromisoformat(e["incurred_on"]), start, end)
    ]


def transactions_between(transactions: list, start: date, end: date, direction=None) -> list:
    out = []
    for t in transactions:
        if direction is not None and t.get("direction") != direction:
            continue
        if in_range(date.fromisoformat(t["occurred_on"]), start, end):
            out.append(t)
    return out


# ── aggregations shared across reports ───────────────────────────────────────
def revenue(sales: list) -> Decimal:
    """Gross takings (sum of receipt totals)."""
    return sum((D(s["total"]) for s in sales), ZERO)


def cogs(sales: list) -> Decimal:
    """Cost of goods sold — the snapshot production cost of what went out."""
    return sum((D(s["cost_of_goods"]) for s in sales), ZERO)


def product_sales(sales: list) -> dict:
    """Per-product roll-up of sold line items.

    Returns ``{product_id: {name, sku, units, revenue, cost}}`` where ``cost``
    is the snapshot COGS of those units (``unit_cost × quantity``).
    """
    out: dict[str, dict] = {}
    for s in sales:
        for it in s["items"]:
            pid = it["product"]
            row = out.setdefault(pid, {
                "name": it["product_name"], "sku": it.get("product_sku", ""),
                "units": ZERO, "revenue": ZERO, "cost": ZERO,
            })
            qty = D(it["quantity"])
            row["units"] += qty
            row["revenue"] += D(it["line_total"])
            row["cost"] += D(it["unit_cost"]) * qty
    return out


def product_production(runs: list) -> dict:
    """Per-product roll-up of completed production runs.

    Returns ``{product_id: {name, units, cost, runs}}``.
    """
    out: dict[str, dict] = {}
    for r in runs:
        pid = r["product"]
        row = out.setdefault(pid, {
            "name": r.get("product_name", "(unknown)"),
            "units": ZERO, "cost": ZERO, "runs": 0,
        })
        row["units"] += D(r["quantity"])
        row["cost"] += D(r["cost"])
        row["runs"] += 1
    return out


def payment_breakdown(sales: list) -> dict:
    """``{payment_method: {revenue, receipts}}`` for the given sales."""
    out: dict[str, dict] = {}
    for s in sales:
        row = out.setdefault(s["payment_method"], {"revenue": ZERO, "receipts": 0})
        row["revenue"] += D(s["total"])
        row["receipts"] += 1
    return out


def expense_total(expenses: list, expense_txns: list) -> Decimal:
    """Money spent: recorded expenses plus manual expense-direction transactions."""
    total = sum((D(e["amount"]) for e in expenses), ZERO)
    total += sum((D(t["amount"]) for t in expense_txns), ZERO)
    return total


#: Expense category used by (now-disabled) automatic inventory write-offs. Such
#: rows are a *non-cash* accounting move — wasted stock, no money left the till —
#: so the reports keep them out of expense / cash totals and surface them
#: separately. Historic data still carries these rows; new waste creates none.
INVENTORY_WRITE_OFF_CATEGORY = "Inventory write-off"


def is_write_off(expense: dict) -> bool:
    """True when an expense row is an inventory write-off (non-cash waste)."""
    return expense.get("category_name") == INVENTORY_WRITE_OFF_CATEGORY


def split_write_offs(expenses: list) -> tuple[list, list]:
    """Partition expense rows into ``(real_expenses, write_offs)``."""
    real, write_offs = [], []
    for e in expenses:
        (write_offs if is_write_off(e) else real).append(e)
    return real, write_offs


def transaction_total(transactions: list) -> Decimal:
    return sum((D(t["amount"]) for t in transactions), ZERO)


def env_decimal(name: str):
    """Read a money/number env var as a Decimal, or ``None`` if unset/blank."""
    raw = os.environ.get(name, "").strip()
    if not raw:
        return None
    try:
        return Decimal(raw)
    except InvalidOperation:
        return None


def resolve_optional(cli_value, env_name: str):
    """CLI value wins; otherwise fall back to an env var; otherwise ``None``."""
    if cli_value is not None:
        return D(cli_value)
    return env_decimal(env_name)


# ── output helpers ───────────────────────────────────────────────────────────
def reconfigure_stdio() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass


def rule(char: str = "═", width: int = 70) -> str:
    return char * width


def header(title: str, subtitle: str = "") -> None:
    print(rule())
    print(title)
    if subtitle:
        print(subtitle)
    print(rule())


def section(title: str) -> None:
    print()
    print(title)
    print(rule("─"))


def kv(label: str, value: str, width: int = 26) -> None:
    print(f"  {label:<{width}} {value}")


def table(headers: list[str], rows: list[list[str]], aligns: list[str] | None = None) -> None:
    """Render a simple fixed-width table. ``aligns`` items are 'l' or 'r'."""
    if not rows:
        print("  (none)")
        return
    cols = len(headers)
    aligns = aligns or ["l"] * cols
    widths = [len(h) for h in headers]
    for row in rows:
        for i in range(cols):
            widths[i] = max(widths[i], len(str(row[i])))

    def fmt(cells):
        out = []
        for i, c in enumerate(cells):
            c = str(c)
            out.append(c.rjust(widths[i]) if aligns[i] == "r" else c.ljust(widths[i]))
        return "  " + "  ".join(out)

    print(fmt(headers))
    print("  " + "  ".join("-" * w for w in widths))
    for row in rows:
        print(fmt(row))


def _json_default(o):
    if isinstance(o, Decimal):
        return float(q2(o))
    if isinstance(o, (date, datetime)):
        return o.isoformat()
    return str(o)


def emit(payload: dict) -> None:
    """Print a report's structured payload as JSON (for ``--json`` / cron)."""
    print(json.dumps(payload, indent=2, default=_json_default))


# ── document sinks (terminal + Word) ─────────────────────────────────────────
# Each report's ``render()`` writes to one of these sinks instead of calling
# ``print`` directly. ``TerminalDoc`` reproduces the original fixed-width text
# byte-for-byte (it just forwards to the helpers above); ``DocxDoc`` builds an
# equivalent Word document. Both expose the same vocabulary — ``header``,
# ``section``, ``kv``, ``table`` and ``text`` — so a single ``render`` serves
# both. Call ``save(path)`` at the end (a no-op for the terminal).
class TerminalDoc:
    """Sink that prints the classic fixed-width terminal report."""

    def header(self, title: str, subtitle: str = "") -> None:
        header(title, subtitle)

    def section(self, title: str) -> None:
        section(title)

    def kv(self, label: str, value: str, width: int = 26) -> None:
        kv(label, value, width)

    def table(self, headers, rows, aligns=None) -> None:
        table(headers, rows, aligns)

    def text(self, line: str = "") -> None:
        print(line)

    def save(self, path: str) -> None:  # output already streamed to stdout
        pass


class DocxDoc:
    """Sink that builds a Word (.docx) document mirroring the report layout.

    Consecutive :meth:`kv` calls are grouped into a borderless two-column table
    (label / value); :meth:`table` renders as a bordered table with a bold
    header row and right-aligned numeric columns. Requires ``python-docx``
    (``pip install python-docx``); the dependency is only imported here, so the
    text and JSON outputs never need it.
    """

    def __init__(self, currency: str = CURRENCY):
        try:
            from docx import Document
        except ModuleNotFoundError:
            sys.exit(
                "The --docx option needs the python-docx package.\n"
                "Install it with:  pip install python-docx"
            )
        self.currency = currency
        self._doc = Document()
        self._kv = None  # the two-column table currently collecting kv() rows

    @staticmethod
    def _bold(cell) -> None:
        for run in cell.paragraphs[0].runs:
            run.bold = True

    def _flush_kv(self) -> None:
        self._kv = None

    def header(self, title: str, subtitle: str = "") -> None:
        self._flush_kv()
        self._doc.add_heading(title.title(), level=0)
        if subtitle:
            p = self._doc.add_paragraph(subtitle)
            if p.runs:
                p.runs[0].italic = True

    def section(self, title: str) -> None:
        self._flush_kv()
        self._doc.add_heading(title, level=1)

    def kv(self, label: str, value: str, width: int = 26) -> None:
        if self._kv is None:
            self._kv = self._doc.add_table(rows=0, cols=2)
        cells = self._kv.add_row().cells
        cells[0].text = label.strip()
        cells[1].text = str(value)
        self._bold(cells[0])

    def table(self, headers, rows, aligns=None) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        self._flush_kv()
        if not rows:
            self._doc.add_paragraph("(none)")
            return
        aligns = aligns or ["l"] * len(headers)
        t = self._doc.add_table(rows=1, cols=len(headers))
        try:
            t.style = "Light Grid Accent 1"
        except KeyError:
            t.style = "Table Grid"
        for i, h in enumerate(headers):
            t.rows[0].cells[i].text = h
            self._bold(t.rows[0].cells[i])
        for row in rows:
            cells = t.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = str(value)
                if aligns[i] == "r":
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    def text(self, line: str = "") -> None:
        # Blank lines are spacing in the terminal; Word handles that itself.
        if not line.strip():
            return
        self._flush_kv()
        stripped = line.strip()
        p = self._doc.add_paragraph(stripped)
        if stripped.endswith(":") and p.runs:  # a "Foo:" sub-label
            p.runs[0].bold = True

    def save(self, path: str) -> None:
        self._flush_kv()
        self._doc.save(path)


def base_arg_parser(description: str) -> argparse.ArgumentParser:
    """An argument parser pre-seeded with the connection + output flags every
    report shares."""
    p = argparse.ArgumentParser(
        description=description,
        epilog=ENV_HELP,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    p.add_argument("--base-url", default=DEFAULT_BASE_URL)
    p.add_argument("--email", default=DEFAULT_EMAIL)
    p.add_argument("--password", default=DEFAULT_PASSWORD)
    p.add_argument("--currency", default=CURRENCY, help="Currency code shown next to amounts.")
    p.add_argument("--json", action="store_true", help="Emit machine-readable JSON instead of text.")
    p.add_argument("--docx", metavar="PATH", default=None,
                   help="Write the report to a Word .docx file at PATH (needs python-docx).")
    return p
