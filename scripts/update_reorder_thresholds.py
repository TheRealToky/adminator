#!/usr/bin/env python3
r"""Bulk-update raw-material reorder thresholds from an Excel file.

Reads ``reorder_thresholds.xlsx`` and, for every row, sets the
``reorder_threshold`` of the matching raw material via the REST API
(``PATCH /catalog/raw-materials/{id}/``). This is the "Reorder threshold"
column shown on the raw-materials page; the app raises a low-stock alert when
on-hand quantity falls to or below it.

Sheet layout (first row is a header, matched case-insensitively; column order
does not matter):

    SKU            | Name (optional)   | Reorder Threshold
    RM-FLOUR-001   | Bread flour       | 5000
    RM-SUGAR-001   | Caster sugar      | 2000

Matching
--------
Rows are matched to a raw material by **SKU** first (unique, exact — after
trimming and upper-casing). If the SKU cell is blank, the script falls back to
an exact case/accent-insensitive **Name** match. Anything it can't resolve
confidently — or a name that matches more than one material — is reported and
skipped rather than guessed. The ``Name`` column is optional and only used for
the fallback / to make the sheet readable.

Rows whose threshold is blank are skipped. A threshold equal to the material's
current value is a no-op (reported as "unchanged", no write). Values must be a
non-negative number; the field stores 2 decimal places.

Safe by default: runs in **dry-run** mode (matches + planned changes only, no
writes). Pass ``--commit`` to actually apply them.

Usage
-----
    # preview what would change (no writes):
    python scripts/update_reorder_thresholds.py

    # actually apply the new thresholds:
    python scripts/update_reorder_thresholds.py --commit

    # point at a different file / stack:
    python scripts/update_reorder_thresholds.py --commit \
        --file "C:\path\reorder_thresholds.xlsx" \
        --base-url http://localhost:8001/api/v1

    # write a ready-to-fill example spreadsheet next to the script and exit:
    python scripts/update_reorder_thresholds.py --make-example

Requirements (host): ``pip install openpyxl`` (only extra dependency; HTTP uses
the standard library).
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import unicodedata
import urllib.error
import urllib.request
from decimal import Decimal, InvalidOperation

try:
    import openpyxl
except ModuleNotFoundError:
    sys.exit('openpyxl is not installed.\n  Fix: python -m pip install openpyxl')

# ── defaults (override on the CLI or via env) ────────────────────────────────
DEFAULT_BASE_URL = os.environ.get("ADMINATOR_API", "http://localhost:8001/api/v1")
DEFAULT_EMAIL = os.environ.get("ADMINATOR_EMAIL", "admin@adminator.local")
DEFAULT_PASSWORD = os.environ.get("ADMINATOR_PASSWORD", "admin12345")
DEFAULT_FILE = os.environ.get(
    "REORDER_XLSX",
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "reorder_thresholds.xlsx"),
)

# Recognised header labels (normalized) -> logical column.
HEADER_ALIASES = {
    "sku": "sku",
    "name": "name",
    "raw material": "name",
    "material": "name",
    "reorder threshold": "threshold",
    "reorder": "threshold",
    "threshold": "threshold",
    "reorder point": "threshold",
    "reorder level": "threshold",
}

# quantize thresholds to the field's precision (DecimalField 2 dp)
TWO_DP = Decimal("0.01")


# ── small helpers ────────────────────────────────────────────────────────────
def normalize(name) -> str:
    """Lowercase, strip accents, collapse to space-separated alphanumeric tokens."""
    if name is None:
        return ""
    s = unicodedata.normalize("NFKD", str(name))
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^a-z0-9]+", " ", s.lower())
    return re.sub(r"\s+", " ", s).strip()


def parse_threshold(raw):
    """Return a non-negative Decimal (2 dp) or raise ValueError with a reason."""
    if raw is None or str(raw).strip() == "":
        raise ValueError("blank")
    try:
        value = Decimal(str(raw).strip())
    except (InvalidOperation, ValueError):
        raise ValueError(f"not a number: {raw!r}")
    if value < 0:
        raise ValueError(f"negative: {value}")
    return value.quantize(TWO_DP)


# ── API client (stdlib urllib; JWT bearer) ───────────────────────────────────
class ApiError(Exception):
    def __init__(self, status: int, body):
        self.status = status
        self.body = body
        super().__init__(f"HTTP {status}: {body}")

    @property
    def detail(self) -> str:
        if isinstance(self.body, dict):
            d = self.body.get("detail")
            if d:
                return str(d)
            return json.dumps(self.body)
        return str(self.body)


class Api:
    def __init__(self, base_url: str, email: str, password: str):
        self.base = base_url.rstrip("/")
        self.token: str | None = None
        _, data = self._request(
            "POST", "/auth/login/", {"email": email, "password": password}, auth=False
        )
        self.token = data["access"]

    def _request(self, method, path, payload=None, auth=True):
        url = self.base + path
        body = json.dumps(payload).encode() if payload is not None else None
        req = urllib.request.Request(url, data=body, method=method)
        req.add_header("Content-Type", "application/json")
        if auth and self.token:
            req.add_header("Authorization", f"Bearer {self.token}")
        try:
            with urllib.request.urlopen(req, timeout=60) as resp:
                raw = resp.read().decode()
                return resp.status, (json.loads(raw) if raw else None)
        except urllib.error.HTTPError as exc:
            raw = exc.read().decode()
            try:
                parsed = json.loads(raw)
            except ValueError:
                parsed = {"detail": raw[:300]}
            raise ApiError(exc.code, parsed)
        except urllib.error.URLError as exc:
            sys.exit(f"Could not reach {url}: {exc}. Is the scratch stack up?")

    def get_all(self, path: str) -> list:
        """Follow DRF pagination and return every result row."""
        sep = "&" if "?" in path else "?"
        _, data = self._request("GET", f"{path}{sep}page_size=1000")
        if isinstance(data, dict) and "results" in data:
            rows = list(data["results"])
            while data.get("next"):
                nxt = data["next"].split("/api/v1", 1)[-1]
                _, data = self._request("GET", nxt)
                rows.extend(data["results"])
            return rows
        return data or []

    def set_reorder_threshold(self, material_id, threshold):
        _, data = self._request(
            "PATCH", f"/catalog/raw-materials/{material_id}/",
            {"reorder_threshold": str(threshold)},
        )
        return data


# ── material index / matching ────────────────────────────────────────────────
def index_materials(materials: list[dict]):
    """Build SKU and name lookups. Name maps to a *list* to flag duplicates."""
    by_sku: dict[str, dict] = {}
    by_name: dict[str, list[dict]] = {}
    for m in materials:
        by_sku[str(m["sku"]).strip().upper()] = m
        by_name.setdefault(normalize(m["name"]), []).append(m)
    return by_sku, by_name


def resolve(row, by_sku, by_name):
    """Return (material|None, reason). reason explains a None result."""
    sku = (row.get("sku") or "").strip().upper()
    if sku:
        m = by_sku.get(sku)
        if m:
            return m, "sku"
        return None, f"no raw material with SKU {sku!r}"
    name = row.get("name") or ""
    if not normalize(name):
        return None, "no SKU and no name"
    cands = by_name.get(normalize(name), [])
    if not cands:
        return None, f"no raw material named {name!r}"
    if len(cands) > 1:
        skus = ", ".join(sorted(c["sku"] for c in cands))
        return None, f"name {name!r} is ambiguous (SKUs: {skus}) — add a SKU"
    return cands[0], "name"


# ── spreadsheet ──────────────────────────────────────────────────────────────
def read_rows(path: str) -> list[dict]:
    """Read the sheet into dicts keyed by logical column (sku/name/threshold)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)

    try:
        header = next(rows_iter)
    except StopIteration:
        sys.exit(f"{path} is empty.")

    col_map: dict[int, str] = {}
    for idx, cell in enumerate(header):
        logical = HEADER_ALIASES.get(normalize(cell))
        if logical and logical not in col_map.values():
            col_map[idx] = logical
    if "threshold" not in col_map.values():
        sys.exit(
            "Could not find a 'Reorder Threshold' column in the header.\n"
            f"  Header seen: {[c for c in header]}\n"
            "  Tip: run with --make-example to generate a correctly-shaped file."
        )

    out = []
    for i, r in enumerate(rows_iter, start=2):
        record = {"row": i, "sku": "", "name": "", "threshold": None}
        for idx, logical in col_map.items():
            value = r[idx] if idx < len(r) else None
            record[logical] = "" if value is None else str(value).strip() if logical != "threshold" else value
        # skip fully-blank rows
        if not (record["sku"] or record["name"]) and record["threshold"] in (None, ""):
            continue
        out.append(record)
    return out


def write_example(path: str) -> None:
    """Write a small, ready-to-fill example workbook."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reorder Thresholds"
    ws.append(["SKU", "Name", "Reorder Threshold"])
    sample = [
        ("RM-FLOUR-001", "Bread flour", 5000),
        ("RM-SUGAR-001", "Caster sugar", 2000),
        ("RM-BUTTER-001", "Butter", 1500),
        ("RM-EGG-001", "Eggs", 120),
        ("RM-YEAST-001", "Fresh yeast", 300),
    ]
    for sku, name, threshold in sample:
        ws.append([sku, name, threshold])
    widths = {"A": 16, "B": 22, "C": 18}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    wb.save(path)


# ── main ─────────────────────────────────────────────────────────────────────
def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass

    p = argparse.ArgumentParser(
        description="Update raw-material reorder thresholds from an Excel file."
    )
    p.add_argument("--file", default=DEFAULT_FILE, help="Path to reorder_thresholds.xlsx")
    p.add_argument("--base-url", default=DEFAULT_BASE_URL)
    p.add_argument("--email", default=DEFAULT_EMAIL)
    p.add_argument("--password", default=DEFAULT_PASSWORD)
    p.add_argument("--commit", action="store_true", help="Actually write (default: dry-run preview).")
    p.add_argument("--make-example", action="store_true",
                   help="Write an example spreadsheet to --file and exit (no network).")
    args = p.parse_args()

    if args.make_example:
        write_example(args.file)
        print(f"Wrote example spreadsheet → {args.file}")
        return 0

    mode = "COMMIT" if args.commit else "DRY-RUN (no writes)"
    print(f"Reorder-threshold update · {mode}")
    print(f"File   : {args.file}")
    print(f"Target : {args.base_url}\n")

    rows = read_rows(args.file)

    api = Api(args.base_url, args.email, args.password)
    materials = api.get_all("/catalog/raw-materials/")
    by_sku, by_name = index_materials(materials)
    print(f"Loaded {len(materials)} raw materials.\n")

    planned, unchanged, skipped, unmatched, bad_value = [], [], [], [], []

    for r in rows:
        # validate the threshold first — a bad value is worth reporting even if
        # the material can't be matched.
        try:
            new_threshold = parse_threshold(r["threshold"])
        except ValueError as exc:
            label = r["sku"] or r["name"] or f"row {r['row']}"
            if str(exc) == "blank":
                skipped.append((r, "blank threshold"))
            else:
                bad_value.append((r, str(exc)))
                print(f"  ! BAD VALUE  {label:<24} — {exc} — skipped")
            continue

        material, reason = resolve(r, by_sku, by_name)
        if material is None:
            unmatched.append((r, reason))
            label = r["sku"] or r["name"] or f"row {r['row']}"
            print(f"  ? UNMATCHED  {label:<24} — {reason} — skipped")
            continue

        current = Decimal(str(material["reorder_threshold"])).quantize(TWO_DP)
        if current == new_threshold:
            unchanged.append((r, material))
            continue

        print(f"  • {material['name']:<26} [{material['sku']}]  "
              f"{current} → {new_threshold}  (by {reason})")
        planned.append((material, current, new_threshold))

    done, failed = [], []
    if args.commit and planned:
        print("\nApplying…")
        for material, current, new_threshold in planned:
            try:
                api.set_reorder_threshold(material["id"], new_threshold)
                done.append((material, new_threshold))
                print(f"  ✓ {material['name']} [{material['sku']}]  → {new_threshold}")
            except ApiError as exc:
                failed.append((material, exc))
                print(f"  ✗ {material['name']} [{material['sku']}]  — {exc.detail}")

    # ── summary ──
    print("\n" + "=" * 60)
    print("Summary")
    print(f"  changes planned      : {len(planned)}")
    print(f"  already up to date   : {len(unchanged)}")
    print(f"  skipped (blank)      : {len(skipped)}")
    print(f"  unmatched            : {len(unmatched)}"
          + (f"  → {', '.join((x[0]['sku'] or x[0]['name']) for x in unmatched)}" if unmatched else ""))
    print(f"  bad values           : {len(bad_value)}"
          + (f"  → {', '.join((x[0]['sku'] or x[0]['name']) for x in bad_value)}" if bad_value else ""))
    if args.commit:
        print(f"  applied              : {len(done)}")
        print(f"  failed               : {len(failed)}")
    else:
        print("\n  DRY-RUN — nothing written. Re-run with --commit to apply these changes.")
    return 1 if failed or bad_value or unmatched else 0


if __name__ == "__main__":
    sys.exit(main())
