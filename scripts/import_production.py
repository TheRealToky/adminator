#!/usr/bin/env python3
r"""Import a day's production sheet into Adminator (production runs).

Reads ``pastry_production.xlsx`` and, for every row, records a production run
of column **F ("Production")** units of the product in column **A ("Product")**
via the REST API (``POST /production/runs/execute/``). The API does the real
work: it consumes the recipe's raw materials, draws down processed materials,
and adds the finished units to stock.

Edge case handled
-----------------
*If a production run is unsuccessful because a raw material is short*, the
script tops that material up (``POST /inventory/stock/receive/`` for the exact
deficit) and retries — looping until the run goes through. (Product runs can
only ever fail on a *raw* material: processed-material consumption is allowed
to go negative server-side, so it never blocks a run.)

Rows with no production quantity (blank or 0) are skipped — there is nothing to
bake. Product names on the sheet are matched to the catalog case/accent-
insensitively, with a small curated alias table for the French⇄English and
word-order differences; anything that can't be matched confidently is reported
and skipped rather than guessed.

Safe by default: runs in **dry-run** mode (matches + plan only, no writes).
Pass ``--commit`` to actually record the runs.

Usage
-----
    # preview what would happen (no writes):
    python scripts/import_production.py

    # actually record the runs:
    python scripts/import_production.py --commit

    # point at a different file / stack / date:
    python scripts/import_production.py --commit \
        --file "C:\path\pastry_production.xlsx" \
        --base-url http://localhost:8001/api/v1 \
        --date 2026-06-16

Requirements (host): ``pip install openpyxl`` (only extra dependency; HTTP uses
the standard library).
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import unicodedata
import urllib.error
import urllib.request
from decimal import ROUND_CEILING, Decimal
from difflib import SequenceMatcher

try:
    import openpyxl
except ModuleNotFoundError:
    sys.exit('openpyxl is not installed.\n  Fix: python -m pip install openpyxl')

# ── defaults (override on the CLI or via env) ────────────────────────────────
DEFAULT_BASE_URL = os.environ.get("ADMINATOR_API", "http://localhost:8001/api/v1")
DEFAULT_EMAIL = os.environ.get("ADMINATOR_EMAIL", "admin@adminator.local")
DEFAULT_PASSWORD = os.environ.get("ADMINATOR_PASSWORD", "admin12345")
DEFAULT_FILE = os.environ.get(
    "PRODUCTION_XLSX",
    r"D:\PC DISAINE\toky\Perso-D\red\production\2026_07\2026_07_20\pastry_production.xlsx",
)
DEFAULT_DATE = "2026-07-20"

# Sheet/column layout: A=Product, F=Production qty (0-indexed 0 and 5).
SHEET_NAME = "Production"
COL_PRODUCT = 0
COL_PRODUCTION = 5
COL_NOTES = 6

# Curated aliases: production-sheet name (normalized) -> exact catalog name.
# Only the cases plain normalization + fuzzy matching can't resolve on their
# own (translations, regional names, word order). Edit here if a name is wrong.
PRODUCT_ALIASES = {
    "baguette au levain": "Sourdough baguette",
    "baguette traditionnelle": "Traditional Baguette",
    "pain au levain": "Sourdough bread",
    "pain complet": "Brown bread",
    "pain aux grains": "Baguette with seeds",
    "cinnamon rolls": "Cinnamon roll",
    "bon muffins": "Muffin",
    "mangue passion": "Tarte mangue passion",
    "tarte chocolat": "Tarte au chocolat",
    "moka cafe": "Mocha",
    "quiche au thon": "Quiche",
}

MATCH_THRESHOLD = 0.82      # fuzzy score below which we refuse to guess
MAX_TOPUP_ROUNDS = 60       # safety cap on raw-material top-up / retry loop

INSUFFICIENT_RE = re.compile(
    r"Not enough stock of (.+?): have ([\d.]+), requested ([\d.]+)\."
)


# ── small helpers ────────────────────────────────────────────────────────────
def normalize(name: str) -> str:
    """Lowercase, strip accents, reduce to space-separated alphanumeric tokens.

    Parenthetical content is *kept* (as tokens) so distinct catalog entries like
    "Cake" and "Cake (6 persons)" don't normalize to the same string.
    """
    if name is None:
        return ""
    s = unicodedata.normalize("NFKD", str(name))
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.lower()
    s = re.sub(r"[^a-z0-9]+", " ", s)              # punctuation/() -> space
    return re.sub(r"\s+", " ", s).strip()


def parse_insufficient(detail: str | None):
    """Return (item_name, have, requested) from an InsufficientStock message."""
    if not detail:
        return None
    m = INSUFFICIENT_RE.search(detail)
    if not m:
        return None
    return m.group(1), Decimal(m.group(2)), Decimal(m.group(3))


def topup_amount(have: Decimal, requested: Decimal) -> Decimal:
    """Smallest receivable quantity that covers the shortfall (4 dp, rounded up)."""
    deficit = Decimal(requested) - Decimal(have)
    if deficit <= 0:
        deficit = Decimal("0.0001")
    return deficit.quantize(Decimal("0.0001"), rounding=ROUND_CEILING)


# ── API client (stdlib urllib; JWT bearer) ───────────────────────────────────
class ApiError(Exception):
    def __init__(self, status: int, body):
        self.status = status
        self.body = body
        super().__init__(f"HTTP {status}: {body}")

    @property
    def detail(self) -> str:
        if isinstance(self.body, dict):
            d = self.body.get("detail")
            if d:
                return str(d)
            return json.dumps(self.body)
        return str(self.body)


class Api:
    def __init__(self, base_url: str, email: str, password: str):
        self.base = base_url.rstrip("/")
        self.token: str | None = None
        _, data = self._request(
            "POST", "/auth/login/", {"email": email, "password": password}, auth=False
        )
        self.token = data["access"]

    def _request(self, method, path, payload=None, auth=True):
        url = self.base + path
        body = json.dumps(payload).encode() if payload is not None else None
        req = urllib.request.Request(url, data=body, method=method)
        req.add_header("Content-Type", "application/json")
        if auth and self.token:
            req.add_header("Authorization", f"Bearer {self.token}")
        try:
            with urllib.request.urlopen(req, timeout=60) as resp:
                raw = resp.read().decode()
                return resp.status, (json.loads(raw) if raw else None)
        except urllib.error.HTTPError as exc:
            raw = exc.read().decode()
            try:
                parsed = json.loads(raw)
            except ValueError:
                parsed = {"detail": raw[:300]}
            raise ApiError(exc.code, parsed)
        except urllib.error.URLError as exc:
            sys.exit(f"Could not reach {url}: {exc}. Is the scratch stack up?")

    def get_all(self, path: str) -> list:
        """Follow DRF pagination and return every result row."""
        sep = "&" if "?" in path else "?"
        _, data = self._request("GET", f"{path}{sep}page_size=1000")
        if isinstance(data, dict) and "results" in data:
            rows = list(data["results"])
            while data.get("next"):
                nxt = data["next"].split("/api/v1", 1)[-1]
                _, data = self._request("GET", nxt)
                rows.extend(data["results"])
            return rows
        return data or []

    def execute_production(self, product_id, quantity, scheduled_for, notes=""):
        _, data = self._request(
            "POST", "/production/runs/execute/",
            {"product": product_id, "quantity": str(quantity),
             "scheduled_for": scheduled_for, "status": "completed", "notes": notes},
        )
        return data

    def receive_material(self, raw_material_id, quantity, note="", reference=""):
        _, data = self._request(
            "POST", "/inventory/stock/receive/",
            {"raw_material": raw_material_id, "quantity": str(quantity),
             "note": note, "reference": reference},
        )
        return data


# ── name matching ────────────────────────────────────────────────────────────
class Matcher:
    """Resolve a sheet product name to a catalog product."""

    def __init__(self, products: list[dict], aliases: dict[str, str]):
        self.products = products
        self.by_norm: dict[str, dict] = {}
        for p in products:
            self.by_norm.setdefault(normalize(p["name"]), p)
        # Resolve alias *targets* to products up front; warn on any typo.
        self.alias_norm: dict[str, dict] = {}
        for sheet_name, catalog_name in aliases.items():
            target = self.by_norm.get(normalize(catalog_name))
            if target is None:
                print(f"  ! alias target not in catalog: {catalog_name!r}", file=sys.stderr)
            else:
                self.alias_norm[normalize(sheet_name)] = target

    def match(self, name: str):
        """Return (product|None, method, score)."""
        n = normalize(name)
        if not n:
            return None, "empty", 0.0
        if n in self.alias_norm:
            return self.alias_norm[n], "alias", 1.0
        if n in self.by_norm:
            return self.by_norm[n], "exact", 1.0
        # Fuzzy: combine straight ratio with an order-independent token ratio so
        # "Baguette Brown" still matches "Brown baguette".
        best, best_score = None, 0.0
        n_sorted = " ".join(sorted(n.split()))
        for p in self.products:
            pn = normalize(p["name"])
            score = max(
                SequenceMatcher(None, n, pn).ratio(),
                SequenceMatcher(None, n_sorted, " ".join(sorted(pn.split()))).ratio(),
            )
            if score > best_score:
                best, best_score = p, score
        if best_score >= MATCH_THRESHOLD:
            return best, "fuzzy", best_score
        return None, "unmatched", best_score


def index_materials(stock_rows: list[dict]) -> dict[str, list[dict]]:
    """Map normalized raw-material name -> [{id, name, stock}] for top-ups."""
    idx: dict[str, list[dict]] = {}
    for row in stock_rows:
        if row.get("kind") != "raw_material" or not row.get("raw_material"):
            continue
        entry = {"id": row["raw_material"], "name": row["item_name"],
                 "stock": Decimal(str(row["quantity"]))}
        idx.setdefault(normalize(row["item_name"]), []).append(entry)
    return idx


def resolve_material(materials: dict[str, list[dict]], name: str, have: Decimal):
    """Find the raw material named in an error; disambiguate dupes by on-hand qty."""
    cands = materials.get(normalize(name), [])
    if not cands:
        return None
    if len(cands) == 1:
        return cands[0]
    for c in cands:                       # duplicate names: match the short one
        if c["stock"] == have:
            return c
    return cands[0]


# ── core: record one run, auto-fixing raw-material shortfalls ─────────────────
def produce_with_topup(api: Api, product: dict, quantity: Decimal, date: str,
                       materials: dict, log) -> dict:
    """Record a production run, topping up any short raw material and retrying."""
    notes = f"Imported from pastry_production.xlsx ({date})"
    last_short = None
    for _ in range(MAX_TOPUP_ROUNDS):
        try:
            return api.execute_production(product["id"], quantity, date, notes)
        except ApiError as exc:
            if exc.status != 400:
                raise
            info = parse_insufficient(exc.detail)
            if info is None:
                raise
            mat_name, have, requested = info
            mat = resolve_material(materials, mat_name, have)
            if mat is None:
                raise RuntimeError(
                    f"raw material {mat_name!r} (from error) not found in catalog"
                )
            amount = topup_amount(have, requested)
            api.receive_material(
                mat["id"], amount,
                note=f"Auto top-up so '{product['name']}' production can run (import)",
                reference="IMPORT-PROD",
            )
            mat["stock"] = have + amount       # keep local view current
            log(f"      ↑ topped up {mat_name}: +{amount} (had {have}, needed {requested})")
            last_short = mat_name
    raise RuntimeError(
        f"gave up after {MAX_TOPUP_ROUNDS} top-up rounds (stuck on {last_short})"
    )


# ── spreadsheet ──────────────────────────────────────────────────────────────
def read_rows(path: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == 1:
            continue  # header
        name = r[COL_PRODUCT] if len(r) > COL_PRODUCT else None
        if name is None or str(name).strip() == "":
            continue
        qty = r[COL_PRODUCTION] if len(r) > COL_PRODUCTION else None
        note = r[COL_NOTES] if len(r) > COL_NOTES else None
        rows.append({"row": i, "name": str(name).strip(), "qty": qty, "note": note})
    return rows


# ── main ─────────────────────────────────────────────────────────────────────
def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass

    p = argparse.ArgumentParser(description="Import production runs from the pastry production xlsx.")
    p.add_argument("--file", default=DEFAULT_FILE, help="Path to pastry_production.xlsx")
    p.add_argument("--base-url", default=DEFAULT_BASE_URL)
    p.add_argument("--email", default=DEFAULT_EMAIL)
    p.add_argument("--password", default=DEFAULT_PASSWORD)
    p.add_argument("--date", default=DEFAULT_DATE, help="scheduled_for date (YYYY-MM-DD)")
    p.add_argument("--commit", action="store_true", help="Actually write (default: dry-run preview).")
    p.add_argument("--limit", type=int, default=0, help="Only process the first N producible rows (testing).")
    args = p.parse_args()

    mode = "COMMIT" if args.commit else "DRY-RUN (no writes)"
    print(f"Production import · {mode}")
    print(f"File   : {args.file}")
    print(f"Target : {args.base_url}  (date {args.date})\n")

    rows = read_rows(args.file)

    api = Api(args.base_url, args.email, args.password)
    products = api.get_all("/catalog/products/")
    materials = index_materials(api.get_all("/inventory/stock/?kind=raw_material"))
    matcher = Matcher(products, PRODUCT_ALIASES)
    print(f"Loaded {len(products)} products, {sum(len(v) for v in materials.values())} stocked materials.\n")

    planned, skipped_zero, unmatched, done, failed = [], [], [], [], []

    for r in rows:
        qty_raw = r["qty"]
        try:
            qty = Decimal(str(qty_raw)) if qty_raw is not None else Decimal("0")
        except Exception:
            qty = Decimal("0")
        if qty <= 0:
            skipped_zero.append(r)
            continue
        product, method, score = matcher.match(r["name"])
        if product is None:
            unmatched.append((r, score))
            print(f"  ? UNMATCHED  {r['name']!r}  (best score {score:.2f}) — skipped")
            continue
        tag = method if method != "fuzzy" else f"fuzzy {score:.2f}"
        print(f"  • {r['name']:<26} → {product['name']:<24} ×{qty}  [{tag}]")
        planned.append((r, product, qty))

    if args.limit:
        planned = planned[: args.limit]

    if args.commit and planned:
        print("\nRecording runs…")
        for r, product, qty in planned:
            try:
                run = produce_with_topup(api, product, qty, args.date, materials,
                                         log=print)
                done.append((product, qty, run))
                print(f"  ✓ {product['name']} ×{qty}  (run {run['id'][:8]}, cost {run['cost']})")
            except (ApiError, RuntimeError) as exc:
                failed.append((product, qty, exc))
                msg = exc.detail if isinstance(exc, ApiError) else str(exc)
                print(f"  ✗ {product['name']} ×{qty}  — {msg}")

    # ── summary ──
    print("\n" + "=" * 60)
    print("Summary")
    print(f"  producible rows planned : {len(planned)}")
    print(f"  skipped (no qty / 0)    : {len(skipped_zero)}"
          + (f"  → {', '.join(x['name'] for x in skipped_zero)}" if skipped_zero else ""))
    print(f"  unmatched (skipped)     : {len(unmatched)}"
          + (f"  → {', '.join(x[0]['name'] for x in unmatched)}" if unmatched else ""))
    if args.commit:
        print(f"  runs recorded           : {len(done)}")
        print(f"  runs failed             : {len(failed)}")
    else:
        print("\n  DRY-RUN — nothing written. Re-run with --commit to record these runs.")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
